"""Bounded file extraction for text and multimodal search."""
from dataclasses import dataclass
from pathlib import Path
import zipfile
import xml.etree.ElementTree as ET
import warnings
import shutil
import subprocess
import json
import math
import wave

VIDEO_EXTENSIONS = set('.mp4 .mkv .mov .avi .webm .m4v .mpeg .mpg .wmv .3gp .mts .m2ts'.split())
AUDIO_EXTENSIONS = set('.wav .mp3 .flac .ogg .opus .m4a .aac .aiff .aif .wma .amr'.split())
from uuid import uuid4

IMAGE_EXTENSIONS = set('.png .jpg .jpeg .gif .webp .bmp .tif .tiff .ico .ppm .pgm .avif .heic'.split())


@dataclass
class Chunk:
    kind: str
    value: str
    location: str
    snippet: str


TEXT_EXTENSIONS = set('.txt .md .rst .csv .tsv .json .jsonl .yaml .yml .toml .ini .cfg .conf .xml .html .htm .css .js .jsx .ts .tsx .py .sh .bash .zsh .c .h .cpp .hpp .cc .java .go .rs .rb .php .sql .log .tex .svg .r .nix .fish .lua .pl .swift .kt'.split())


def extract(path: Path, scratch: Path) -> list[Chunk]:
    """Extract chunks; the caller owns scratch and its lifetime."""
    path, scratch = Path(path), Path(scratch)
    if path.suffix.lower() in {'.docx', '.xlsx', '.pptx', '.odt', '.ods', '.odp'}:
        try:
            with zipfile.ZipFile(path) as archive:
                if sum(i.file_size for i in archive.infolist()) > 64 * 1024 * 1024:
                    raise ValueError(f'{path.name}: office decompressed size exceeds 64 MB; split the document')
                for entry in archive.infolist():
                    if entry.filename.lower().endswith('.xml'):
                        xml = archive.read(entry).replace(b'\x00', b'').upper()
                        if b'<!DOCTYPE' in xml or b'<!ENTITY' in xml:
                            raise ValueError(f'{path.name}: unsafe XML entities/DTD; re-export the document without entities')
            return _extract(path, scratch)
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError(f'{path.name}: corrupt or invalid office document; repair or re-export it ({exc})') from exc
    if path.suffix.lower() in AUDIO_EXTENSIONS | VIDEO_EXTENSIONS:
        probe = json.loads(_run(['ffprobe', '-v', 'error', '-show_format', '-show_streams', '-of', 'json', str(path.resolve())], path))
        try:
            duration = float(probe['format']['duration'])
            if not math.isfinite(duration) or duration <= 0:
                raise ValueError('non-positive or non-finite duration')
        except (KeyError, TypeError, ValueError) as exc:
            raise ValueError(f'{path.name}: invalid or unknown media duration; repair or re-export the recording') from exc
        if duration > 7200:
            raise ValueError(f'{path.name}: media exceeds 2 hours; split the recording before indexing')
        chunks = []
        if path.suffix.lower() in VIDEO_EXTENSIONS:
            for start in range(0, math.ceil(duration), 5):
                output = (scratch / f'{uuid4().hex}.png').resolve()
                _run(['ffmpeg', '-v', 'error', '-nostdin', '-y', '-ss', str(start), '-i', str(path.resolve()), '-map', '0:v:0', '-frames:v', '1', '-vf', 'scale=1024:1024:force_original_aspect_ratio=decrease', str(output)], path)
                chunks.append(_image(output, scratch, f'video {start}s', path.name))
                output.unlink()
        has_audio = any(s.get('codec_type') == 'audio' for s in probe.get('streams', []))
        if path.suffix.lower() in AUDIO_EXTENSIONS and not has_audio:
            raise ValueError(f'{path.name}: missing audio stream; repair or re-export the recording')
        if has_audio:
            chunks.extend(_audio(path, scratch, duration))
        return chunks
    if path.suffix.lower() in IMAGE_EXTENSIONS:
        return [_image(path, scratch, 'image')]
    if path.suffix.lower() == '.pdf':
        try:
            return _extract(path, scratch)
        except Exception as exc:
            raise ValueError(f'{path.name}: corrupt or invalid PDF; repair or re-export it ({exc})') from exc
    return _extract(path, scratch)


def _audio(path: Path, scratch: Path, duration: float) -> list[Chunk]:
    chunks = []
    for start in range(0, math.ceil(duration), 20):
        output = (scratch / f'{uuid4().hex}.wav').resolve()
        _run(['ffmpeg', '-v', 'error', '-nostdin', '-y', '-ss', str(start), '-i', str(path.resolve()), '-t', str(min(20, duration - start)), '-map', '0:a:0', '-vn', '-ac', '1', '-ar', '16000', '-c:a', 'pcm_s16le', str(output)], path)
        with wave.open(str(output)) as audio:
            if audio.getnframes() == 0:
                raise ValueError(f'{path.name}: empty decoded audio at {start}s; repair or re-export media')
        location = f'audio {start}-{min(start + 20, duration):g}s'
        chunks.append(Chunk('audio', str(output), location, f'{path.name}: {location}'))
    return chunks


def _run(args: list[str], path: Path) -> str:
    try:
        return subprocess.run(args, check=True, capture_output=True, text=True, timeout=120).stdout
    except subprocess.TimeoutExpired as exc:
        raise ValueError(f'{path.name}: extraction timed out after 120 seconds; split the file') from exc
    except (OSError, subprocess.CalledProcessError) as exc:
        detail = getattr(exc, 'stderr', '') or str(exc)
        raise ValueError(f'{path.name}: media/render extraction failed; verify the file and required tools ({detail[-1000:]})') from exc


def _image(path: Path, scratch: Path, location: str, original_name: str | None = None) -> Chunk:
    from PIL import Image, ImageOps
    try:
        with warnings.catch_warnings():
            warnings.simplefilter('error', Image.DecompressionBombWarning)
            with Image.open(path) as source:
                source.verify()
            with Image.open(path) as source:
                source.draft('RGB', (1024, 1024))
                source.thumbnail((1024, 1024))
                image = ImageOps.exif_transpose(source).convert('RGB')
                output = (scratch / f'{uuid4().hex}.png').resolve()
                image.save(output, 'PNG')
        return Chunk('image', str(output), location, f'{original_name or path.name}: {location}')
    except Exception as exc:
        raise ValueError(f'{path.name}: corrupt, invalid or excessive-pixel image; repair or resize it ({exc})') from exc


def _extract(path: Path, scratch: Path) -> list[Chunk]:
    if path.suffix.lower() == '.pdf':
        from pypdf import PdfReader
        chunks = []
        for i, page in enumerate(PdfReader(path).pages, 1):
            text = page.extract_text() or ''
            if text.strip():
                chunks.extend(_text_chunks(text, f'page {i}'))
            else:
                renderer = shutil.which('pdftoppm')
                if not renderer:
                    raise ValueError(f'{path.name}: page {i} has no text; install pdftoppm (Poppler) to index scanned pages')
                prefix = (scratch / uuid4().hex).resolve()
                _run([renderer, '-f', str(i), '-l', str(i), '-singlefile', '-scale-to', '1024', '-png', str(path.resolve()), str(prefix)], path)
                rendered = prefix.with_suffix('.png')
                chunks.append(_image(rendered, scratch, f'page {i}', path.name))
                rendered.unlink()
        return chunks
    if path.suffix.lower() in {'.odt', '.ods', '.odp'}:
        with zipfile.ZipFile(path) as archive:
            root = ET.fromstring(archive.read('content.xml'))
        ns = {'text': 'urn:oasis:names:tc:opendocument:xmlns:text:1.0', 'table': 'urn:oasis:names:tc:opendocument:xmlns:table:1.0', 'draw': 'urn:oasis:names:tc:opendocument:xmlns:drawing:1.0'}
        if path.suffix.lower() == '.ods':
            parts = [(e, f'sheet {e.get("{" + ns["table"] + "}name", "unnamed")}') for e in root.findall('.//table:table', ns)]
        elif path.suffix.lower() == '.odp':
            parts = [(e, f'slide {i}') for i, e in enumerate(root.findall('.//draw:page', ns), 1)]
        else:
            parts = [(root, 'document')]
        return [c for e, location in parts for c in _text_chunks('\n'.join(''.join(p.itertext()) for p in e.iter() if p.tag in {'{' + ns['text'] + '}p', '{' + ns['text'] + '}h'}), location)]
    if path.suffix.lower() == '.pptx':
        from pptx import Presentation
        chunks = []
        for number, slide in enumerate(Presentation(path).slides, 1):
            text = []
            for shape in _shapes(slide.shapes):
                if shape.has_text_frame:
                    text.append(shape.text)
                if shape.has_table:
                    text.extend('\t'.join(c.text for c in row.cells) for row in shape.table.rows)
            chunks.extend(_text_chunks('\n'.join(text), f'slide {number}'))
        return chunks
    if path.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            return [chunk for sheet in book for chunk in _text_chunks('\n'.join('\t'.join('' if v is None else str(v) for v in row) for row in sheet.iter_rows(values_only=True)), f'sheet {sheet.title}')]
        finally:
            book.close()
    if path.suffix.lower() == '.docx':
        from docx import Document
        document = Document(path)
        text = '\n'.join([p.text for p in document.paragraphs] + ['\t'.join(c.text for c in row.cells) for table in document.tables for row in table.rows])
        return _text_chunks(text, 'document')
    if path.stat().st_size > 32 * 1024 * 1024:
        if path.suffix.lower() not in TEXT_EXTENSIONS:
            with path.open('rb') as stream:
                sample = stream.read(8192)
            try:
                _decode_text(sample)
            except UnicodeError:
                return _metadata(path)
        raise ValueError(f'{path.name}: text exceeds 32 MB; split the file before indexing')
    try:
        text = _decode_text(path.read_bytes())
    except UnicodeError as exc:
        if path.suffix.lower() in TEXT_EXTENSIONS:
            raise ValueError(f'{path.name}: invalid UTF-8/UTF-16 text; convert the encoding before indexing') from exc
        return _metadata(path)
    return _text_chunks(text, 'text')


def _metadata(path: Path) -> list[Chunk]:
    metadata = f'Unsupported content: filename {path.name}; type {path.suffix or "unknown"}; size {path.stat().st_size} bytes'
    return [Chunk('text', metadata, 'metadata only', metadata)]


def _decode_text(data: bytes) -> str:
    encoding = 'utf-8-sig'
    if data.startswith((b'\xff\xfe', b'\xfe\xff')):
        encoding = 'utf-16'
    elif b'\x00' in data and len(data) % 2 == 0 and max(data[1::2].count(0), data[::2].count(0)) > len(data) // 4:
        encoding = 'utf-16-le' if data[1::2].count(0) > data[::2].count(0) else 'utf-16-be'
    text = data.decode(encoding)
    if any(ord(c) < 32 and c not in '\n\r\t\f' for c in text):
        raise UnicodeError('binary control characters')
    return text


def _shapes(shapes):
    for shape in shapes:
        yield shape
        if hasattr(shape, 'shapes'):
            yield from _shapes(shape.shapes)


def _text_chunks(text: str, location: str) -> list[Chunk]:
    chunks = []
    for start in range(0, len(text), 1650):
        part = text[start:start + 1800]
        chunks.append(Chunk('text', part, f'{location}, chars {start + 1}-{start + len(part)}', part[:240]))
        if start + 1800 >= len(text):
            break
    return chunks
