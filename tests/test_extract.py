from pathlib import Path
import pytest
import zipfile
import shutil
import subprocess
import wave
import struct
import json


def test_pptx_grouped_text(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    from search_tool.extract import extract
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    box = group.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    box.text = 'Grouped content'
    source = tmp_path / 'group.pptx'
    deck.save(source)
    assert 'Grouped content' in '\n'.join(c.value for c in extract(source, tmp_path))


def test_audio_file_without_audio_stream_refused(tmp_path, monkeypatch):
    from search_tool.extract import extract
    source = tmp_path / 'empty.wav'
    source.write_bytes(b'placeholder')
    payload = {'format': {'duration': '1'}, 'streams': []}
    monkeypatch.setattr(subprocess, 'run', lambda args, **kw: subprocess.CompletedProcess(args, 0, stdout=json.dumps(payload)))
    with pytest.raises(ValueError, match='audio stream.*re-export'):
        extract(source, tmp_path)


def test_office_xml_entities_refused(tmp_path):
    from search_tool.extract import extract
    source = tmp_path / 'entities.odt'
    with zipfile.ZipFile(source, 'w') as archive:
        archive.writestr('content.xml', '<!DOCTYPE x [<!ENTITY a "expanded">]><x>&a;</x>')
    with pytest.raises(ValueError, match='entities.*XML'):
        extract(source, tmp_path)


def test_invalid_office_xml_valueerror_has_filename(tmp_path):
    from search_tool.extract import extract
    source = tmp_path / 'broken.docx'
    with zipfile.ZipFile(source, 'w') as archive:
        archive.writestr('[Content_Types].xml', '<invalid>')
    with pytest.raises(ValueError, match='broken.docx.*repair'):
        extract(source, tmp_path)


def test_large_unsupported_binary_metadata_not_text_limit(tmp_path):
    from search_tool.extract import extract
    source = tmp_path / 'large.bin'
    with source.open('wb') as f:
        f.write(b'\xffbinary')
        f.truncate(33 * 1024 * 1024)
    c = extract(source, tmp_path)[0]
    assert c.location == 'metadata only' and str(source.stat().st_size) in c.value


def test_large_extensionless_text_refused(tmp_path):
    from search_tool.extract import extract
    source = tmp_path / 'README'
    with source.open('wb') as f:
        f.write(b'hello ' * 6000000)
    with pytest.raises(ValueError, match='32.*MB'):
        extract(source, tmp_path)


@pytest.mark.parametrize('duration', [None, 'NaN', '-1', '0', 'Infinity'])
def test_unusable_media_duration_actionable(tmp_path, monkeypatch, duration):
    from search_tool.extract import extract
    source = tmp_path / 'unknown.wav'
    source.write_bytes(b'placeholder')
    payload = {'format': {} if duration is None else {'duration': duration}, 'streams': [{'codec_type': 'audio'}]}
    def probe(args, **kwargs):
        assert args[0] == 'ffprobe', 'invalid duration must be rejected before decoding'
        return subprocess.CompletedProcess(args, 0, stdout=json.dumps(payload))
    monkeypatch.setattr(subprocess, 'run', probe)
    with pytest.raises(ValueError, match='duration.*re-export'):
        extract(source, tmp_path)


def test_typescript_is_text_not_video(tmp_path):
    from search_tool.extract import extract
    source = tmp_path / 'app.ts'
    source.write_text('export const answer: number = 42;')
    assert extract(source, tmp_path)[0].value == source.read_text()



@pytest.mark.parametrize('extension', ['wav', 'mp4'])
def test_corrupt_media_actionable(tmp_path, extension):
    from search_tool.extract import extract
    source = tmp_path / f'broken.{extension}'
    source.write_bytes(b'not valid media')
    with pytest.raises(ValueError, match='broken.*failed|broken.*invalid'):
        extract(source, tmp_path)


def test_media_over_two_hours_refused(tmp_path, monkeypatch):
    if not shutil.which('ffprobe'):
        pytest.skip('ffprobe not installed')
    from search_tool.extract import extract
    source = tmp_path / 'long.wav'
    data_size = 7201 * 16000 * 2
    header = b'RIFF' + struct.pack('<I', 36 + data_size) + b'WAVEfmt ' + struct.pack('<IHHIIHH', 16, 1, 1, 16000, 32000, 2, 16) + b'data' + struct.pack('<I', data_size)
    with source.open('wb') as f:
        f.write(header)
        f.truncate(44 + data_size)
    real_run = subprocess.run
    def refuse_decode(args, **kwargs):
        if args[0] == 'ffmpeg':
            pytest.fail('oversized media reached decoder instead of being refused')
        return real_run(args, **kwargs)
    monkeypatch.setattr(subprocess, 'run', refuse_decode)
    with pytest.raises(ValueError, match='2 hours'):
        extract(source, tmp_path)


def test_media_subprocess_timeout_actionable(tmp_path, monkeypatch):
    from search_tool.extract import extract
    source = tmp_path / 'slow.wav'
    source.write_bytes(b'placeholder')
    def timeout(args, **kwargs):
        assert isinstance(args, list)
        assert kwargs['timeout'] > 0 and not kwargs.get('shell', False)
        raise subprocess.TimeoutExpired(args, kwargs['timeout'])
    monkeypatch.setattr(subprocess, 'run', timeout)
    with pytest.raises(ValueError, match='timed out.*split'):
        extract(source, tmp_path)


@pytest.mark.parametrize('audio', [True, False])
def test_video_five_second_frames_plus_audio(tmp_path, audio):
    from PIL import Image
    from search_tool.extract import extract
    source = make_media(tmp_path, video=True, audio=audio, duration=11)
    chunks = extract(source, tmp_path)
    images = [c for c in chunks if c.kind == 'image']
    sounds = [c for c in chunks if c.kind == 'audio']
    assert len(images) == 3
    assert all(source.name in c.snippet for c in images)
    assert [c.location for c in images] == ['video 0s', 'video 5s', 'video 10s']
    assert len(sounds) == int(audio)
    for c in images:
        with Image.open(c.value) as image:
            assert image.mode == 'RGB' and max(image.size) <= 1024


def make_media(tmp_path, video=False, audio=True, duration=21):
    if not shutil.which('ffmpeg') or not shutil.which('ffprobe'):
        pytest.skip('ffmpeg/ffprobe not installed')
    source = tmp_path / ('clip.mp4' if video else 'tone.wav')
    args = ['ffmpeg', '-v', 'error', '-y']
    if video:
        args += ['-f', 'lavfi', '-i', 'color=c=blue:s=128x72:r=10']
    if audio:
        args += ['-f', 'lavfi', '-i', 'sine=frequency=440:sample_rate=44100']
    args += ['-t', str(duration)]
    if video:
        args += ['-c:v', 'mpeg4']
        if audio:
            args += ['-c:a', 'aac']
    args += [str(source)]
    subprocess.run(args, check=True, capture_output=True, timeout=30)
    return source


def test_audio_decoded_mono_16k_twenty_second_chunks(tmp_path):
    from search_tool.extract import extract
    source = make_media(tmp_path)
    chunks = extract(source, tmp_path)
    assert len(chunks) == 2 and all(c.kind == 'audio' for c in chunks)
    assert '0' in chunks[0].location and '20' in chunks[1].location
    frames = []
    for c in chunks:
        assert Path(c.value).is_absolute()
        with wave.open(c.value) as audio:
            assert audio.getnchannels() == 1
            assert audio.getframerate() == 16000
            assert audio.getsampwidth() == 2
            assert 0 < audio.getnframes() <= 20 * 16000
            frames.append(audio.getnframes())
    assert sum(frames) == 21 * 16000


def test_scanned_pdf_missing_renderer_actionable(tmp_path, monkeypatch):
    from pypdf import PdfWriter
    from search_tool.extract import extract
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    source = tmp_path / 'scan.pdf'
    writer.write(source)
    monkeypatch.setattr(shutil, 'which', lambda _: None)
    with pytest.raises(ValueError, match='pdftoppm'):
        extract(source, tmp_path)


def test_scanned_pdf_rendered_when_available(tmp_path):
    if not shutil.which('pdftoppm'):
        pytest.skip('pdftoppm not installed')
    from pypdf import PdfWriter
    from PIL import Image
    from search_tool.extract import extract
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=200)
    source = tmp_path / 'scan.pdf'
    writer.write(source)
    chunks = extract(source, tmp_path)
    assert len(chunks) == 1 and chunks[0].kind == 'image'
    assert 'page 1' in chunks[0].location
    assert source.name in chunks[0].snippet
    with Image.open(chunks[0].value) as image:
        assert image.mode == 'RGB' and max(image.size) <= 1024


def test_image_normalized_bounded_rgb_png(tmp_path):
    from PIL import Image
    from search_tool.extract import extract
    source = tmp_path / 'picture.jpg'
    Image.new('L', (2048, 1200), 120).save(source)
    scratch = tmp_path / 'scratch'
    scratch.mkdir()
    chunks = extract(source, scratch)
    assert len(chunks) == 1 and chunks[0].kind == 'image'
    output = Path(chunks[0].value)
    assert output.is_absolute() and output.parent == scratch
    with Image.open(output) as image:
        assert image.mode == 'RGB' and image.format == 'PNG'
        assert image.size == (1024, 600)


@pytest.mark.parametrize('extension', ['jpg', 'png', 'pdf'])
def test_corrupt_image_pdf_actionable(tmp_path, extension):
    from search_tool.extract import extract
    source = tmp_path / f'broken.{extension}'
    source.write_bytes(b'not valid content')
    with pytest.raises(ValueError, match='broken.*corrupt|broken.*invalid'):
        extract(source, tmp_path)


def test_image_decompression_bomb_refused(tmp_path, monkeypatch):
    from PIL import Image
    from search_tool.extract import extract
    source = tmp_path / 'bomb.png'
    Image.new('RGB', (20, 20)).save(source)
    monkeypatch.setattr(Image, 'MAX_IMAGE_PIXELS', 250)
    with pytest.raises(ValueError, match='bomb.*pixel|bomb.*decompression'):
        extract(source, tmp_path)


def test_pdf_per_page_text(tmp_path):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    from search_tool.extract import extract
    writer = PdfWriter()
    for text in ['First page', 'Second page']:
        page = writer.add_blank_page(width=300, height=200)
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
        stream = DecodedStreamObject()
        stream.set_data(f'BT /F1 12 Tf 10 100 Td ({text}) Tj ET'.encode())
        page[NameObject('/Contents')] = writer._add_object(stream)
    source = tmp_path / 'pages.pdf'
    writer.write(source)
    chunks = extract(source, tmp_path)
    assert any('page 1' in c.location and 'First page' in c.value for c in chunks)
    assert any('page 2' in c.location and 'Second page' in c.value for c in chunks)


@pytest.mark.parametrize('extension', ['docx', 'xlsx', 'pptx', 'odt', 'ods', 'odp'])
def test_corrupt_office_actionable(tmp_path, extension):
    from search_tool.extract import extract
    source = tmp_path / f'broken.{extension}'
    source.write_bytes(b'not an office archive')
    with pytest.raises(ValueError, match='broken.*corrupt|broken.*invalid'):
        extract(source, tmp_path)


@pytest.mark.parametrize('extension', ['docx', 'xlsx', 'pptx', 'odt', 'ods', 'odp'])
def test_office_archive_size_limit(tmp_path, extension):
    from search_tool.extract import extract
    source = tmp_path / f'bomb.{extension}'
    with zipfile.ZipFile(source, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('content.xml', b'x' * (64 * 1024 * 1024 + 1))
    with pytest.raises(ValueError, match='64.*MB'):
        extract(source, tmp_path)


@pytest.mark.parametrize('extension,body,location', [
    ('odt', '<office:text><text:p>Hello <text:span>ODT</text:span></text:p></office:text>', 'document'),
    ('ods', '<office:spreadsheet><table:table table:name="Budget"><table:table-row><table:table-cell><text:p>Revenue</text:p></table:table-cell></table:table-row></table:table></office:spreadsheet>', 'Budget'),
    ('odp', '<office:presentation><draw:page draw:name="Intro"><draw:frame><text:p>Launch</text:p></draw:frame></draw:page></office:presentation>', 'slide 1'),
])
def test_open_document_xml(tmp_path, extension, body, location):
    from search_tool.extract import extract
    source = tmp_path / f'office.{extension}'
    xml = '<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0"><office:body>' + body + '</office:body></office:document-content>'
    with zipfile.ZipFile(source, 'w') as archive:
        archive.writestr('content.xml', xml)
    chunks = extract(source, tmp_path)
    assert any(location in c.location for c in chunks)
    expected = {'odt': 'Hello ODT', 'ods': 'Revenue', 'odp': 'Launch'}[extension]
    assert expected in '\n'.join(c.value for c in chunks)


def test_pptx_slide_text_tables(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    from search_tool.extract import extract
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = 'Launch plan'
    slide.shapes.add_table(1, 1, Inches(1), Inches(2), Inches(3), Inches(1)).table.cell(0, 0).text = 'Milestone'
    deck.slides.add_slide(deck.slide_layouts[1]).shapes.title.text = 'Risks'
    source = tmp_path / 'deck.pptx'
    deck.save(source)
    chunks = extract(source, tmp_path)
    assert any('slide 1' in c.location and 'Launch plan' in c.value and 'Milestone' in c.value for c in chunks)
    assert any('slide 2' in c.location and 'Risks' in c.value for c in chunks)


def test_xlsx_each_sheet(tmp_path):
    from openpyxl import Workbook
    from search_tool.extract import extract
    book = Workbook()
    book.active.title = 'Budget'
    book.active.append(['Revenue', 123])
    book.create_sheet('People').append(['Ada', 'Engineer'])
    source = tmp_path / 'book.xlsx'
    book.save(source)
    chunks = extract(source, tmp_path)
    assert any('Budget' in c.location and 'Revenue' in c.value and '123' in c.value for c in chunks)
    assert any('People' in c.location and 'Ada' in c.value for c in chunks)


def test_docx_paragraphs_tables(tmp_path):
    from docx import Document
    from search_tool.extract import extract
    doc = Document()
    doc.add_paragraph('Project alpha')
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Owner'
    table.cell(0, 1).text = 'Ada'
    source = tmp_path / 'report.docx'
    doc.save(source)
    chunks = extract(source, tmp_path)
    text = '\n'.join(c.value for c in chunks)
    assert 'Project alpha' in text and 'Owner' in text and 'Ada' in text
    assert all(c.kind == 'text' for c in chunks)


def test_text_over_32mb_refused(tmp_path):
    from search_tool.extract import extract
    source = tmp_path / 'huge.txt'
    with source.open('wb') as f:
        f.truncate(32 * 1024 * 1024 + 1)
    with pytest.raises(ValueError, match='32.*MB'):
        extract(source, tmp_path)


def test_corrupt_utf8_text_refused(tmp_path):
    from search_tool.extract import extract
    source = tmp_path / 'broken.py'
    source.write_bytes(b'\xffbroken')
    with pytest.raises(ValueError, match='UTF'):
        extract(source, tmp_path)


@pytest.mark.parametrize('data', [b'\x89BIN\xff\x00\x01', b'\x00\x01\x02\x03'])
def test_unknown_binary_explicit_metadata(tmp_path, data):
    from search_tool.extract import extract
    source = tmp_path / 'unknown.bin'
    source.write_bytes(data)
    chunks = extract(source, tmp_path)
    assert len(chunks) == 1
    c = chunks[0]
    assert c.kind == 'text' and c.location == 'metadata only'
    assert 'unsupported content' in c.snippet.lower()
    assert source.name in c.value and '.bin' in c.value and str(len(data)) in c.value


def test_text_chunk_overlap_preserves_all_content(tmp_path):
    from search_tool.extract import extract
    text = ''.join(f'{i:06d} café ' for i in range(600))
    source = tmp_path / 'code.py'
    source.write_text(text)
    chunks = extract(source, tmp_path)
    assert len(chunks) > 1
    assert all(0 < len(c.value) <= 1800 for c in chunks)
    assert all(a.value[-150:] == b.value[:150] for a, b in zip(chunks, chunks[1:]))
    assert chunks[0].value + ''.join(c.value[150:] for c in chunks[1:]) == text


@pytest.mark.parametrize('encoding', ['utf-16', 'utf-16-le', 'utf-16-be'])
def test_extensionless_utf16(tmp_path, encoding):
    from search_tool.extract import extract
    source = tmp_path / 'README'
    source.write_bytes('Hello café\ncode = 12'.encode(encoding))
    assert extract(source, tmp_path)[0].value == 'Hello café\ncode = 12'


def test_utf8_text(tmp_path):
    from search_tool.extract import Chunk, extract
    source = tmp_path / 'notes.txt'
    source.write_text('Hello café world', encoding='utf-8')
    result = extract(source, tmp_path)
    assert len(result) == 1
    assert isinstance(result[0], Chunk)
    assert result[0].kind == 'text'
    assert result[0].value == 'Hello café world'
    assert result[0].location
    assert 'café' in result[0].snippet
